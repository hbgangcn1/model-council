"""Offline tests for render_formats (no live API/LLM calls)."""

import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from orchestrator.render_formats import convert_report, extract_title  # noqa: E402

ORCH = Path(__file__).resolve().parent
RUNS = ORCH.parent / "runs"


def _find_real_report() -> Path:
    cands = sorted(RUNS.glob("*/report.md"))
    for c in cands:
        if c.stat().st_size > 500:
            return c
    raise unittest.SkipTest("no real runs/*/report.md fixture found")


class TestRenderFormats(unittest.TestCase):
    def test_convert_real_report_offline(self) -> None:
        src = _find_real_report()
        md_text = src.read_text(encoding="utf-8")
        title = extract_title(md_text, "report")
        self.assertTrue(title and len(title) >= 2)
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp) / "report.md"
            shutil.copyfile(src, work)
            written = convert_report(work, True, True, True, True)
            html_p = work.parent / "report.html"
            docx_p = work.parent / "report.docx"
            pptx_p = work.parent / "report.pptx"
            pdf_p = work.parent / "report.pdf"
            self.assertEqual(
                {p.name for p in written},
                {"report.html", "report.docx", "report.pptx", "report.pdf"},
            )
            self.assertTrue(html_p.is_file())
            self.assertTrue(docx_p.is_file())
            html_text = html_p.read_text(encoding="utf-8")
            # Title text (first heading, markdown markers stripped) intact.
            first_word = re.sub(r"\*+|`+", "", title).split()[0]
            self.assertIn(first_word, html_text)
            self.assertIn("<table", html_text)  # real reports contain tables
            from docx import Document

            doc = Document(str(docx_p))
            self.assertGreaterEqual(len(doc.paragraphs), 1)
            self.assertTrue(any(p.text.strip() for p in doc.paragraphs))
            from pptx import Presentation

            prs = Presentation(str(pptx_p))
            self.assertGreaterEqual(len(prs.slides), 2)
            self.assertTrue(pdf_p.is_file())
            self.assertGreater(pdf_p.stat().st_size, 10 * 1024)

    def test_cli_flags(self) -> None:
        src = _find_real_report()
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp) / "report.md"
            shutil.copyfile(src, work)
            r = subprocess.run(
                [sys.executable, str(ORCH / "render_formats.py"), str(work), "--html"],
                capture_output=True,
                text=True,
            )
            self.assertEqual(r.returncode, 0, r.stderr)
            self.assertTrue((Path(tmp) / "report.html").is_file())
            self.assertFalse((Path(tmp) / "report.docx").exists())


if __name__ == "__main__":
    unittest.main()
